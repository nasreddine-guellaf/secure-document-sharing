from io import BytesIO
from pathlib import Path


SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


class UnsupportedFileTypeError(ValueError):
    pass


def get_file_extension(filename: str) -> str:
    return Path(filename or "").suffix.lower()


def validate_supported_file(filename: str) -> str:
    extension = get_file_extension(filename)

    if extension not in SUPPORTED_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{extension or 'unknown'}'. Supported types: {allowed}."
        )

    return extension


def extract_text_from_txt(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode("utf-8").strip()
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1", errors="ignore").strip()


def extract_text_from_pdf(file_bytes: bytes) -> str:
    import fitz

    pages = []
    with fitz.open(stream=file_bytes, filetype="pdf") as document:
        for page in document:
            pages.append(page.get_text("text"))

    return "\n".join(pages).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(file_bytes))
    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text:
            parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)

    return "\n".join(parts).strip()


def extract_text(filename: str, file_bytes: bytes) -> str:
    extension = validate_supported_file(filename)

    if extension == ".txt":
        return extract_text_from_txt(file_bytes)

    if extension == ".pdf":
        return extract_text_from_pdf(file_bytes)

    if extension == ".docx":
        return extract_text_from_docx(file_bytes)

    raise UnsupportedFileTypeError("Unsupported file type.")
